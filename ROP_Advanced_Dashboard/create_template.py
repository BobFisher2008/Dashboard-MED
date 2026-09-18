from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
import pandas as pd

BASE=Path(__file__).resolve().parent
branches=pd.read_csv(BASE/'data/dim_branch.csv.gz')
branches=branches[branches['Branch_ID']!='NETWORK']

wb=Workbook()
ws=wb.active
ws.title='Inventory_Input'
ws.sheet_view.showGridLines=False

navy='12263A'; teal='1F4E5F'; light='E8F1F2'; blue='DDEBF7'; gray='6B7280'; orange='FCE4D6'
ws.merge_cells('A1:I1')
ws['A1']='Үлдэгдлийн автомат шинэчлэлтийн оролтын загвар'
ws['A1'].font=Font(size=18,bold=True,color='FFFFFF')
ws['A1'].fill=PatternFill('solid',fgColor=navy)
ws['A1'].alignment=Alignment(horizontal='left',vertical='center')
ws.row_dimensions[1].height=32
ws.merge_cells('A2:I2')
ws['A2']='SKU_ID байвал нэрийн mapping шаардахгүй. Branch хоосон бол NETWORK түвшний үлдэгдэл гэж үзнэ.'
ws['A2'].font=Font(size=10,color=gray)
ws['A2'].alignment=Alignment(wrap_text=True)

headers=['SnapshotDate','Branch','SKU_ID','ProductName','Manufacturer','Supplier','OnHand','OnOrder','Backorder']
for c,h in enumerate(headers,1):
    cell=ws.cell(4,c,h)
    cell.font=Font(bold=True,color='FFFFFF')
    cell.fill=PatternFill('solid',fgColor=teal)
    cell.alignment=Alignment(horizontal='center',vertical='center',wrap_text=True)
ws.row_dimensions[4].height=32

# A few blank rows and one illustrative row that users should delete/overwrite.
example=['2026-08-13','NETWORK','', 'Тобрекс нүд/дус 0,3% 5мл N1','','',100,0,0]
for c,v in enumerate(example,1):
    ws.cell(5,c,v)
    ws.cell(5,c).font=Font(color='0000FF')
for r in range(6,105):
    for c in range(1,10):
        ws.cell(r,c,'')
        ws.cell(r,c).font=Font(color='0000FF')

# Add an Excel table so Power Query can ingest the template consistently.
tab=Table(displayName='Inventory_Input',ref='A4:I104')
style=TableStyleInfo(name='TableStyleMedium2',showFirstColumn=False,showLastColumn=False,showRowStripes=True,showColumnStripes=False)
tab.tableStyleInfo=style
ws.add_table(tab)

widths={'A':15,'B':34,'C':12,'D':48,'E':30,'F':30,'G':13,'H':13,'I':13}
for col,w in widths.items(): ws.column_dimensions[col].width=w
ws.freeze_panes='A5'
ws.auto_filter.ref='A4:I104'
for row in ws.iter_rows(min_row=5,max_row=104,min_col=1,max_col=9):
    for cell in row:
        cell.alignment=Alignment(vertical='center',wrap_text=True)
        cell.border=Border(bottom=Side(style='hair',color='D1D5DB'))

# Lists and validation.
list_ws=wb.create_sheet('Lists')
list_ws['A1']='Branch'
list_ws['A2']='NETWORK'
for idx,name in enumerate(branches['BranchName'],3): list_ws.cell(idx,1,name)
list_ws.sheet_state='hidden'
dv=DataValidation(type='list',formula1=f"=Lists!$A$2:$A${len(branches)+2}",allow_blank=True)
ws.add_data_validation(dv)
dv.add('B5:B104')

inst=wb.create_sheet('Заавар')
inst.sheet_view.showGridLines=False
inst.merge_cells('A1:F1');inst['A1']='Ашиглах заавар';inst['A1'].font=Font(size=18,bold=True,color='FFFFFF');inst['A1'].fill=PatternFill('solid',fgColor=navy)
notes=[
('1','ERP-ээс үлдэгдлийн файл экспортлоно. Боломжтой бол SKU_ID болон Branch баганыг заавал оруулна.'),
('2','Inventory_Input хүснэгтэд шинэ мөрүүдийг paste хийнэ. SnapshotDate, OnHand заавал байна.'),
('3','Python refresh: python refresh_inventory.py --inventory <file> --snapshot-date YYYY-MM-DD'),
('4','Power BI refresh хийхэд powerbi_data хавтас дахь CSV-үүд шинэчлэгдсэн байна.'),
('5','OnOrder, Backorder байхгүй бол 0 үлдээнэ. Inventory Position = OnHand + OnOrder - Backorder.'),
('6','Branch хоосон эсвэл NETWORK бол сүлжээний нийт түвшин; салбарын нэр байвал branch-level status гарна.'),
]
inst['A3']='Алхам';inst['B3']='Тайлбар'
for c in ['A3','B3']:
    inst[c].font=Font(bold=True,color='FFFFFF');inst[c].fill=PatternFill('solid',fgColor=teal)
for r,(n,t) in enumerate(notes,4):
    inst.cell(r,1,n);inst.cell(r,2,t);inst.cell(r,2).alignment=Alignment(wrap_text=True,vertical='top')
inst.column_dimensions['A'].width=10;inst.column_dimensions['B'].width=105
inst.freeze_panes='A4'

out=BASE/'Inventory_Upload_Template.xlsx'
wb.save(out)
print(out)
